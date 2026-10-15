import openpyxl
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()
sheet = wb.active
sheet.title = "Semester_Results"

headers = ["Roll No", "Student Name", "Lab Marks (out of 50)", "Theory Marks (out of 50)", "Total Marks (out of 100)", "Result"]
sheet.append(headers)

sample_data = [
    [71, "Pradnya", 45, 40, None, None],
    [72, "Pranshu", 25, 12, None, None],
    [73, "Preksha", 48, 47, None, None],
    [74, "Prisha", 42, 39, None, None],
    [75, "Priyansh", 40, 35, None, None]
]

for row in sample_data:
    sheet.append(row)

total_sum = 0
num_students = 0

for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
    lab_marks = row[2].value
    theory_marks = row[3].value
    
    total_marks = lab_marks + theory_marks
    row[4].value = total_marks
    
    if total_marks >= 40:
        row[5].value = "Pass"
    else:
        row[5].value = "Fail"
        
    total_sum += total_marks
    num_students += 1

header_font = Font(bold=True)
header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

for cell in sheet[1]:
    cell.font = header_font
    cell.fill = header_fill

class_average = total_sum / num_students if num_students > 0 else 0
sheet.append(["", "", "", "Class Average:", class_average, ""])

for cell in sheet[sheet.max_row]:
    cell.font = Font(bold=True)

for col in sheet.columns:
    max_length = 0
    column = col[0].column_letter 
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = (max_length + 2)
    sheet.column_dimensions[column].width = adjusted_width

wb.save("AIML_Batch_Grades.xlsx")